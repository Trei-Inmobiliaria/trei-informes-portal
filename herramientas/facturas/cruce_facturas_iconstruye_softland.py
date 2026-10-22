#!/usr/bin/env python3
"""
Cruce de facturas recibidas IConstruye ↔ Softland (contabilización y pago)
Trei Inmobiliaria / VCB — Gerencia de Administración y Finanzas

Qué hace
--------
1. Extrae de Supabase «BBDD IConstruye» los DTE recibidos (33, 34, 56, 61) de las 10 razones
   sociales del grupo, con su estado en el flujo de compras, aprobadores y estado de pago.
2. Extrae de Supabase «Contabilidad - Finanzas» (esquema grupo, réplica Softland):
   - la contabilización de cada documento (comprobante, cuenta, centro de costo),
   - los pagos por egreso identificados por glosa «Pago: <ttd> <folio>;»,
   - las facturas cedidas a factoring (traspaso 2-1-02-101 → 2-1-02-109) y su pago,
   - el saldo abierto por documento (saldo del auxiliar asignado FIFO a lo más reciente).
3. Cruza ambas fuentes y clasifica cada documento:
   Contabilización: Contabilizada / Pendiente de contabilizar / Provisoria / nula / Sin empresa en Softland
   Pago Softland:   Pagada (egreso) / Pagada a factoring / En Factoring / Cancelada por traspaso /
                    Pagada sin referencia / Pago parcial / En Proveedores / Sin contabilizar / Nota de crédito
4. Agrega el estado del SII (eventos de reclamo/aceptación que carga agente_sii_eventos.py en
   iconstruye.sii_eventos_reclamo): Rechazada / Aceptada / Aceptada (Aut) / Sin pronunciamiento / Sin consulta.
   «Aceptada (Aut)» = sin eventos del SII y vencido el plazo de 8 días corridos desde la recepción (aceptación tácita). El SII es la guía;
   los rechazos manuales del flujo de aprobación de IConstruye NO afectan este estado. El motivo se deduce de la
   regla de aceptación: evento RCD ⇒ «Mal referencia OC-EEPP».
5. Muestra la OC (vía recepciones) y lo que el proveedor referenció en el DTE (código 801 OC, 803 Contrato/EEPP…).
6. Escribe: facturas_cruce.json, Facturas_IConstruye_Softland.xlsx y, si se entrega la plantilla, los dos HTML
   (completo y restringido VCB). Las fechas son dinámicas: por defecto desde el 1-ene del año hasta hoy.

Uso
---
    pip install pandas psycopg2-binary openpyxl
    export IC_DB_URL="postgresql://postgres.jyrbujcnqsemwhkcnwkj:<clave>@<pooler>.pooler.supabase.com:5432/postgres"
    export SF_DB_URL="postgresql://postgres.xdislftwnamrlqqxybtm:<clave>@<pooler>.pooler.supabase.com:5432/postgres"
    python cruce_facturas_iconstruye_softland.py --plantilla plantilla_facturas.html --salida ./salida
    (sin --desde/--hasta/--corte usa: 1-ene del año en curso, hasta mañana (excl.), corte = hoy)

Las cadenas de conexión están en Supabase → Project Settings → Database → Connection string (Session pooler).
No dejar claves escritas en el script ni en archivos compartidos.
"""
import argparse, json, os, re, sys
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo
import pandas as pd

# ------------------------------------------------------------------ parámetros fijos
EMP = {  # RUT receptor → código de empresa en Softland (grupo.empresas)
    '76689694-4': 'IVCB1', '77108223-8': 'IVCB2', '76856039-0': 'IVCB3', '77098757-1': 'IVCB4',
    '76660193-6': 'IVCB5', '77495781-2': 'IVCB13', '77944356-6': 'IVCB14', '77944379-5': 'IVCB25',
    '77532191-1': 'IVCB30', '77911822-3': None}  # Activo Gerenciamiento: no está en la réplica Softland
RS = {'76689694-4': 'VCB Constructora SpA', '77108223-8': 'VCB Inversiones S.A.', '76856039-0': 'Inmobiliaria Q6 SpA',
      '77098757-1': 'Inmobiliaria Q8 SpA', '76660193-6': 'Trei Inmobiliaria SpA', '77495781-2': 'Inmobiliaria Q11 SpA',
      '77944356-6': 'Inmobiliaria Q12 SpA', '77944379-5': 'Inmobiliaria Q14 SpA', '77532191-1': 'Los Boldos SpA',
      '77911822-3': 'Activo Gerenciamiento Inmobiliario SpA'}
TIPO = {'33': 'Factura', '34': 'Factura exenta', '56': 'Nota de débito', '61': 'Nota de crédito'}
TOL = 2  # tolerancia de redondeo en pesos
# Meses sin iConstruye: facturas contabilizadas en Softland con emisión anterior a esta fecha y sin DTE en iConstruye
# se agregan como «Solo Softland» (SII: Aceptada, validadas con el RCV en el punto de control). La constructora (IVCB1) se excluye:
# siempre tuvo iConstruye.
SOLO_SF_HASTA = {'2026': '2026-05-01'}
SOLO_SF_EXCLUYE = {'IVCB1'}
ETQ_SOLO_SF = 'Solo Softland'
SII_HIST = 'Aceptada (Hist)'


def _dv(n):
    s, f = 0, 2
    for d in reversed(str(int(n))):
        s += int(d) * f; f = 2 if f == 7 else f + 1
    r = 11 - s % 11
    return '0' if r == 11 else 'K' if r == 10 else str(r)


def solo_softland(g, ic, anio):
    """Documentos vigentes en Softland (por empresa/tipo/folio/proveedor) sin DTE en iConstruye y emitidos antes del inicio de iConstruye."""
    hasta = SOLO_SF_HASTA.get(str(anio))
    if not hasta or len(g) == 0: return None
    inv = {v: k for k, v in EMP.items() if v}
    ya = set(zip(ic.emp, ic.tipo, ic.folio.astype('Int64'), ic.aux))
    e = g[g.emp.isin(inv) & ~g.emp.isin(SOLO_SF_EXCLUYE) & g.fd.notna() & (pd.to_datetime(g.fd) < pd.Timestamp(hasta))
          & (pd.to_datetime(g.fd) >= pd.Timestamp(f'{anio}-01-01'))].copy()
    e = e[[(a, b, c, d) not in ya for a, b, c, d in zip(e.emp, e.tipo, e.folio.astype('Int64'), e.aux)]]
    if len(e) == 0: return None
    e['tax_id_receptor'] = e.emp.map(inv)
    e['tax_id_emisor'] = [f"{a}-{_dv(a)}" if str(a).isdigit() else str(a) for a in e.aux]
    e['razon_social_emisor'] = e.nom.fillna('')
    e['tipo_dte'] = e.tipo.astype(int)
    e['fe'] = e.fd
    e['monto_total'] = e.m.abs()
    e['tot'] = e.monto_total
    e['solosf'] = True
    e['rs'] = e.tax_id_receptor.map(RS)
    return e[['emp', 'rs', 'aux', 'tipo', 'folio', 'tipo_dte', 'tax_id_emisor', 'razon_social_emisor', 'tax_id_receptor', 'fe',
              'monto_total', 'tot', 'solosf']]

# ------------------------------------------------------------------ SQL — BBDD IConstruye
SQL_IC_DTE = """
with d as (
  select * from iconstruye.dte_chile
  where fecha_emision >= %(desde)s and fecha_emision < %(hasta)s and tipo_dte in (33,34,56,61)
    and tax_id_receptor = any(%(ruts)s)
),
oc as (
  select id_factura_asociada, string_agg(distinct nullif(numero_documento_origen,''), ', ') ocs
  from iconstruye.recepciones where tipo_documento_origen = 'Orden de Compra' and id_factura_asociada is not null group by 1
),
ap as (
  select id_factura_comprador,
         string_agg(coalesce(nombre_usuario_aprobador,id_usuario_aprobador) ||
           case when es_aprobada='SI' then coalesce(' ('||to_char(fecha_aprobacion,'DD-MM-YYYY')||')','') else ' (pendiente)' end,
           ' → ' order by orden_flujo_aprobacion) aprobadores
  from iconstruye.facturas_comprador_aprobaciones group by 1
)
select d.tipo_dte, d.folio, d.tax_id_emisor, d.razon_social_emisor, d.tax_id_receptor,
  d.fecha_emision::date fe, d.fecha_ingreso::date fi, d.fecha_vencimiento::date fv,
  d.monto_neto, d.monto_exento, d.monto_iva, d.monto_total, d.tiene_cesion, d.razon_social_cesionario,
  d.tipo_documento_referencia, d.folio_documento_referencia, d.estado_sii dte_sii, oc.ocs,
  coalesce(f.codigo_centro_gestion, d.codigo_centro_gestion) cg_cod, coalesce(f.nombre_centro_gestion, d.nombre_centro_gestion) cg,
  f.descripcion_estado_factura est_ic, f.descripcion_estado_asociacion asoc, f.descripcion_estado_pago_factura pago,
  f.tipo_aceptacion_reclamo_sii acept, ap.aprobadores,
  f.es_pagada pg, f.fecha_pago::date fp, f.fecha_estimada_pago::date fep
from d
left join iconstruye.facturas_comprador f
  on f.folio_unico = d.folio::text and replace(f.tax_id_proveedor,'.','') = d.tax_id_emisor and f.tax_id_comprador = d.tax_id_receptor
left join ap on ap.id_factura_comprador = f.id_factura_comprador
left join oc on oc.id_factura_asociada = f.id_factura_comprador
order by d.fecha_emision
"""

# Eventos del SII por documento (los carga agente_sii_eventos.py; 'SIN' = consultado sin eventos)
SQL_SII = """
select rut_emisor, tipo_dte, folio, rut_receptor, codigo_evento, descripcion_evento, fecha_evento, fecha_consulta
from iconstruye.sii_eventos_reclamo
"""

# ------------------------------------------------------------------ SQL — Contabilidad (Softland, esquema grupo)
# Contabilización: línea de pasivo 2-1-* del documento + línea de gasto de mayor monto del mismo comprobante
SQL_SF_CONTAB = """
with doc as (
  select m.empresa, m.cpbano, m.cpbnum, m.areacod, m.ttdcod, m.numdoc::bigint numdoc, m.codaux, min(m.cpbfec) cpbfec,
         sum(m.movhaber - m.movdebe) monto_pasivo, min(c.cpbest) cpbest, min(m.movfe) fdoc
  from grupo.movimientos m join grupo.comprobantes c using (empresa,cpbano,cpbnum,areacod)
  where m.ttdcod in ('33','34','56','61') and m.pctcod like '2-1-%%' and m.cpbfec >= %(desde_sf)s
  group by 1,2,3,4,5,6,7
),
gasto as (
  select distinct on (m.empresa,m.cpbano,m.cpbnum,m.areacod) m.empresa,m.cpbano,m.cpbnum,m.areacod, m.pctcod, pc.pcdesc, cc.desccc
  from grupo.movimientos m
  left join grupo.plan_cuentas pc on pc.empresa=m.empresa and pc.pccodi=m.pctcod
  left join grupo.centros_costo cc on cc.empresa=m.empresa and cc.codicc=m.cccod
  where (m.empresa,m.cpbano,m.cpbnum,m.areacod) in (select empresa,cpbano,cpbnum,areacod from doc)
    and m.pctcod not like '2-1-%%' and m.pctcod not like '1-1-07%%' and m.movdebe+m.movhaber>0
  order by m.empresa,m.cpbano,m.cpbnum,m.areacod, abs(m.movdebe-m.movhaber) desc
)
select d.empresa e, d.ttdcod t, d.numdoc n, d.codaux a, d.cpbnum c, d.cpbfec::date f, d.cpbest s, round(d.monto_pasivo) m,
       g.pctcod gc, trim(g.pcdesc) gd, trim(g.desccc) ccd, d.fdoc::date fd, trim(au.nomaux) nomaux
from doc d left join gasto g using (empresa,cpbano,cpbnum,areacod)
left join grupo.auxiliares au on au.empresa = d.empresa and au.codaux = d.codaux
"""

# Pagos identificados por glosa «Pago: <ttd> <folio>;» sobre cuentas 2-1-*, con banco del egreso
SQL_SF_PAGOS = """
with p as (
  select m.empresa, m.cpbano, m.cpbnum, m.areacod, c.cpbtip, m.cpbfec::date f, m.codaux,
         (regexp_match(m.movglosa,'^Pago:\\s*(\\S+)\\s+([0-9]+)'))[1] ttd,
         (regexp_match(m.movglosa,'^Pago:\\s*(\\S+)\\s+([0-9]+)'))[2] nd,
         m.movdebe - m.movhaber monto
  from grupo.movimientos m join grupo.comprobantes c using (empresa,cpbano,cpbnum,areacod)
  where m.pctcod like '2-1-%%' and m.cpbfec >= %(desde)s and c.cpbest='V' and m.movglosa ilike 'pago:%%'
),
banco as (
  select distinct on (m.empresa,m.cpbano,m.cpbnum,m.areacod) m.empresa,m.cpbano,m.cpbnum,m.areacod, trim(pc.pcdesc) banco
  from grupo.movimientos m left join grupo.plan_cuentas pc on pc.empresa=m.empresa and pc.pccodi=m.pctcod
  where (m.empresa,m.cpbano,m.cpbnum,m.areacod) in (select empresa,cpbano,cpbnum,areacod from p where cpbtip='E')
    and m.pctcod like '1-1-0%%' and m.movhaber>0
  order by m.empresa,m.cpbano,m.cpbnum,m.areacod, m.movhaber desc
)
select p.empresa e, p.codaux a, p.ttd t, p.nd n, p.cpbtip k, p.cpbnum c, p.f, round(p.monto) m, b.banco b
from p left join banco b using (empresa,cpbano,cpbnum,areacod)
where p.nd is not null
"""

# Saldo abierto por documento en Proveedores (2-1-02 sin 103 relacionadas ni 109 factoring), año en curso.
# rem = monto − pagos identificados − traspaso a factoring; saldo del auxiliar asignado FIFO a lo más reciente.
SQL_SF_ABIERTO = """
with base as (
  select m.*, c.cpbtip from grupo.movimientos m join grupo.comprobantes c using (empresa,cpbano,cpbnum,areacod)
  where m.cpbano=%(anio)s and c.cpbest='V' and m.pctcod like '2-1-02%%' and m.pctcod not in ('2-1-02-103','2-1-02-109')
),
cpb109 as (select distinct empresa,cpbano,cpbnum,areacod from grupo.movimientos
           where cpbano=%(anio)s and pctcod='2-1-02-109' and movhaber>0 and cpbnum<>'00000000'),
docs as (
  select empresa, codaux, ttdcod, numdoc::bigint numdoc, sum(movhaber) amt, min(movfe) fe
  from base where movhaber>0 and ttdcod not in ('01','02','03','CH','PL','OD','00') group by 1,2,3,4
),
pagos as (
  select empresa, codaux, (regexp_match(movglosa,'^Pago:\\s*(\\S+)\\s+([0-9]+)'))[1] ttd,
         ((regexp_match(movglosa,'^Pago:\\s*(\\S+)\\s+([0-9]+)'))[2])::bigint nd,
         sum(movdebe-movhaber) filter (where cpbtip='E') pe, sum(movdebe-movhaber) filter (where cpbtip<>'E') pt
  from base where movglosa ilike 'pago:%%' group by 1,2,3,4
),
fact as (   -- débito a Proveedores en el comprobante de traspaso a Factoring; folio desde la glosa «FACTORING … <folio>»
  select b.empresa, b.codaux, coalesce((regexp_match(b.movglosa,'([0-9]{2,})'))[1]::bigint, b.numdoc::bigint) nd, sum(b.movdebe) pf
  from base b join cpb109 k using (empresa,cpbano,cpbnum,areacod) where b.movdebe>0 group by 1,2,3
),
saldo as (select empresa, codaux, sum(movhaber-movdebe) s from base group by 1,2),
d2 as (
  select d.*, coalesce(p.pe,0) pe, coalesce(p.pt,0) pt, coalesce(f.pf,0) pf,
         greatest(0, d.amt - coalesce(p.pe,0) - coalesce(p.pt,0) - coalesce(f.pf,0)) rem, s.s
  from docs d
  left join pagos p on p.empresa=d.empresa and p.codaux=d.codaux and p.ttd=d.ttdcod and p.nd=d.numdoc
  left join fact f on f.empresa=d.empresa and f.codaux=d.codaux and f.nd=d.numdoc
  left join saldo s on s.empresa=d.empresa and s.codaux=d.codaux
),
d3 as (select *, sum(rem) over (partition by empresa, codaux order by fe desc, numdoc desc
                                rows between unbounded preceding and 1 preceding) newer from d2)
select empresa e, codaux a, ttdcod t, numdoc n, round(amt) amt, round(pe) pe, round(pt) pt, round(pf) pf,
       round(greatest(0, least(rem, greatest(coalesce(s,0),0) - coalesce(newer,0)))) abierto
from d3 where ttdcod in ('33','34','56')
"""

# Calces de notas de crédito: comprobantes de Proveedores cuya glosa dice «CALCE NCE <folio>» (crédito) y «FACTURA <folio>» (débito)
SQL_SF_CALCE = """
select m.empresa e, m.cpbnum c, m.codaux a, m.movglosa g, round(m.movdebe) d, round(m.movhaber) h
from grupo.movimientos m join grupo.comprobantes c using (empresa,cpbano,cpbnum,areacod)
where m.cpbano=%(anio)s and c.cpbest='V' and m.pctcod like '2-1-02%%'
  and (m.empresa,m.cpbano,m.cpbnum,m.areacod) in (
      select empresa,cpbano,cpbnum,areacod from grupo.movimientos where cpbano=%(anio)s and movglosa ~* '^\\s*calce\\s+n[a-z]*\\s*[0-9]+')
"""

# Movimientos de la cuenta 2-1-02-109 Cuentas por Pagar a Factoring (cesión, refactoring y pago)
SQL_SF_FACTORING = """
select m.empresa e, m.cpbnum c, c.cpbtip k, m.cpbfec::date f, m.pctcod p, m.codaux a, m.ttdcod t, m.numdoc::bigint n,
       round(m.movdebe) d, round(m.movhaber) h, m.movglosa g, trim(au.nomaux) nom
from grupo.movimientos m join grupo.comprobantes c using (empresa,cpbano,cpbnum,areacod)
left join grupo.auxiliares au on au.empresa=m.empresa and au.codaux=m.codaux
where m.cpbano=%(anio)s and c.cpbest='V' and m.pctcod='2-1-02-109'
order by m.empresa, m.cpbfec, m.cpbnum
"""


# ------------------------------------------------------------------ extracción
def query(url, sql, params):
    import psycopg2
    with psycopg2.connect(url) as cn:
        return pd.read_sql_query(sql, cn, params=params)


def extraer(a):
    ic_url, sf_url = os.environ.get('IC_DB_URL'), os.environ.get('SF_DB_URL')
    if not ic_url or not sf_url:
        sys.exit('Faltan las variables de entorno IC_DB_URL y/o SF_DB_URL')
    anio = a.desde[:4]
    desde_sf = f"{int(anio) - 1}-11-01"  # contabilizaciones de ene/feb pueden venir desde nov del año anterior
    p = dict(desde=a.desde, hasta=a.hasta, ruts=list(EMP.keys()), anio=anio, desde_sf=desde_sf)
    print('Extrayendo IConstruye…');  ic = query(ic_url, SQL_IC_DTE, p)
    try: sii = query(ic_url, SQL_SII, p)
    except Exception as e:  # tabla aún no creada: el informe sigue, sin estado SII
        print('  AVISO: no se pudo leer iconstruye.sii_eventos_reclamo:', e); sii = pd.DataFrame()
    print('Extrayendo Softland…')
    sf = query(sf_url, SQL_SF_CONTAB, p)
    pay = query(sf_url, SQL_SF_PAGOS, p)
    opn = query(sf_url, SQL_SF_ABIERTO, p)
    f9 = query(sf_url, SQL_SF_FACTORING, p)
    try: calce = query(sf_url, SQL_SF_CALCE, p)
    except Exception as e:
        print('  AVISO: no se pudieron leer los calces de NC:', e); calce = pd.DataFrame()
    return ic, sf, pay, opn, f9, sii, calce


# ------------------------------------------------------------------ cruce y clasificación
def _v(a):
    if a is None or a is pd.NA or (isinstance(a, float) and pd.isna(a)): return None
    return a


def _d(a):
    a = _v(a)
    if a is None: return None
    s = str(a)[:10]
    return s if s >= '2000' else None  # IConstruye usa 1900-01-01 como «sin fecha»


def _fmt(n):
    return f"{round(n):,}".replace(',', '.')


def _ddmmyyyy(s):
    s = str(s)[:10]; return f"{s[8:10]}-{s[5:7]}-{s[:4]}"


REF_SII = {'801': 'OC', '802': 'Nota de pedido', '803': 'Contrato', '52': 'Guía', '33': 'Factura', '34': 'Factura exenta',
           '56': 'Nota de débito', '61': 'Nota de crédito'}
PLAZO_RECLAMO_DIAS = 8   # días corridos para reclamar un DTE (Ley 19.983)
MOTIVO_RCD = 'Mal referencia OC-EEPP'   # única regla de rechazo vigente (reclamo de contenido, RCD)
ETIQ_RECHAZO = {'RCD': MOTIVO_RCD, 'RFP': 'Reclamo por falta parcial de mercaderías', 'RFT': 'Reclamo por falta total de mercaderías'}


def estados_sii(sii):
    """(rut_emisor, tipo, folio, rut_receptor) -> dict(sii, siim, siif, siic). Gobierna el último evento resolutivo por fecha.
    Rechazada: RCD/RFP/RFT. Aceptada: ACD/ERM, o PAG (pagado al contado) sin otro evento resolutivo → 'Aceptada (Cont)'.
    'SIN' (sentinela) = consultado sin eventos → Sin pronunciamiento."""
    res = {}
    if sii is None or len(sii) == 0: return res
    sii = sii.copy()
    sii['fecha_evento'] = pd.to_datetime(sii.fecha_evento, errors='coerce')
    for k, g in sii.groupby(['rut_emisor', 'tipo_dte', 'folio', 'rut_receptor']):
        g = g.sort_values('fecha_evento', na_position='first')
        cons = g.fecha_consulta.max()
        resol = g[g.codigo_evento.isin(['RCD', 'RFP', 'RFT', 'ACD', 'ERM'])]
        o = dict(siic=str(cons)[:16].replace(' ', 'T') if pd.notna(cons) else None)
        if len(resol):
            u = resol.iloc[-1]
            if u.codigo_evento in ETIQ_RECHAZO:
                o.update(sii='Rechazada', siim=ETIQ_RECHAZO[u.codigo_evento], siif=_d(u.fecha_evento))
            else:
                o.update(sii='Aceptada', siif=_d(u.fecha_evento))
        else:
            pg = g[g.codigo_evento == 'PAG']
            if len(pg):   # «DTE pagado al contado»: el SII no registra aceptación en estos documentos; se tratan como aceptados
                o.update(sii='Aceptada (Cont)', siif=_d(pg.iloc[-1].fecha_evento))
            else:
                o.update(sii='Sin pronunciamiento')
        res[(k[0], int(k[1]), int(k[2]), k[3])] = o
    return res


def enlaces_calce(calce):
    """{(emp, aux, folio_NC): [(folio_factura, monto)]} leído de los calces de Softland (glosas «CALCE NCE n» / «FACTURA n»).
    Dentro de cada comprobante y proveedor se parean primero montos iguales y el resto en orden."""
    res = {}
    if calce is None or len(calce) == 0: return res
    c = calce.copy(); c['aux'] = c.a.astype(str).str.lstrip('0')
    for (e, cp, aux), g in c.groupby(['e', 'c', 'aux']):
        ncs, facs = [], []
        for r in g.itertuples():
            m = re.match(r'^\s*calce\s+n[a-z]*\s*(\d+)', r.g or '', re.I)
            if m and r.h > 0: ncs.append([int(m.group(1)), float(r.h)]); continue
            m = re.match(r'^\s*factura\s+(\d+)', r.g or '', re.I)
            if m and r.d > 0: facs.append([int(m.group(1)), float(r.d)])
        for nc in ncs:                        # montos iguales primero
            for f in facs:
                if f[1] > 0 and abs(f[1] - nc[1]) <= TOL:
                    res.setdefault((e, aux, nc[0]), []).append((f[0], f[1])); f[1] = 0; nc[1] = 0; break
        for nc in ncs:                        # resto en orden
            for f in facs:
                if nc[1] <= TOL: break
                if f[1] > TOL:
                    u = min(nc[1], f[1]); res.setdefault((e, aux, nc[0]), []).append((f[0], u)); nc[1] -= u; f[1] -= u
    return res


def cruzar(ic, sf, pay, opn, f9, corte, sii=None, anio=None, calce=None):
    est_sii = estados_sii(sii)
    # --- IConstruye
    ic = ic.copy()
    ic['emp'] = ic.tax_id_receptor.map(EMP)
    ic['rs'] = ic.tax_id_receptor.map(RS)
    ic['aux'] = ic.tax_id_emisor.str.split('-').str[0].str.lstrip('0')   # codaux Softland = RUT sin DV
    ic['tipo'] = ic.tipo_dte.astype(int).astype(str)
    ic['folio'] = pd.to_numeric(ic.folio).astype('Int64')
    for c in ['monto_neto', 'monto_exento', 'monto_iva', 'monto_total']:
        ic[c] = pd.to_numeric(ic[c])
    ic['tot'] = ic.monto_total.fillna(ic.monto_neto.fillna(0) + ic.monto_exento.fillna(0) + ic.monto_iva.fillna(0))

    # --- Softland: contabilización (llave empresa + tipo + folio + RUT proveedor)
    sf = sf.rename(columns={'e': 'emp', 't': 'tipo', 'n': 'folio'}).copy()
    sf['aux'] = sf.a.astype(str).str.lstrip('0')
    sf['folio'] = pd.to_numeric(sf.folio).astype('Int64')
    sf['tipo'] = sf.tipo.astype(str)
    sf['m'] = pd.to_numeric(sf.m)
    k = ['emp', 'tipo', 'folio', 'aux']
    vig = sf[sf.s == 'V'].sort_values('f')
    if 'fd' not in sf.columns: sf['fd'] = None; sf['nomaux'] = None   # compatibilidad con extracciones sin esas columnas
    vig = sf[sf.s == 'V'].sort_values('f')
    g = vig.groupby(k).agg(c=('c', 'first'), f=('f', 'first'), m=('m', 'sum'), gc=('gc', 'first'),
                           gd=('gd', 'first'), ccd=('ccd', 'first'), fd=('fd', 'min'), nom=('nomaux', 'first')).reset_index()
    extra = solo_softland(g, ic, anio or (str(ic.fe.min())[:4] if len(ic) else '')) if (anio or len(ic)) else None
    if extra is not None:
        ic = pd.concat([ic, extra], ignore_index=True)
    if 'solosf' not in ic.columns: ic['solosf'] = False
    np_ = sf[sf.s != 'V'].groupby(k).agg(s_np=('s', 'first')).reset_index()
    x = ic.merge(g, on=k, how='left').merge(np_, on=k, how='left')

    def est_contab(r):
        if r.emp is None or pd.isna(r.emp): return 'Sin empresa en Softland'
        if pd.notna(r.c): return 'Contabilizada'
        if pd.notna(r.s_np): return 'Provisoria / nula'
        return 'Pendiente de contabilizar'
    x['est_sf'] = x.apply(est_contab, axis=1)

    # --- pagos identificados (detalle de egresos y traspasos)
    pay = pay.copy()
    pay['aux'] = pay.a.astype(str).str.lstrip('0')
    pay['n'] = pd.to_numeric(pay.n, errors='coerce').astype('Int64')
    pay['m'] = pd.to_numeric(pay.m)
    agg = {}
    for r in pay.sort_values('f').itertuples():
        a = agg.setdefault((r.e, r.aux, str(r.t), r.n), {'eg': [], 'tr': []})
        banco = r.b if isinstance(r.b, str) and r.b else 'banco s/i'
        if r.k == 'E': a['eg'].append(f"{r.c} · {_ddmmyyyy(r.f)} · {banco} · {_fmt(r.m)}")
        else: a['tr'].append(f"{r.c} · {_ddmmyyyy(r.f)} · {_fmt(r.m)}")

    # --- saldo abierto por documento
    opn = opn.copy()
    for c in ['amt', 'pe', 'pt', 'pf', 'abierto']: opn[c] = pd.to_numeric(opn[c])
    opn['aux'] = opn.a.astype(str).str.lstrip('0')
    opmap = {(q.e, q.aux, str(q.t), int(q.n)): q for q in opn.itertuples()}

    # --- factoring: saldo por (empresa, folio) en 2-1-02-109
    f9 = f9.copy()
    for c in ['d', 'h']: f9[c] = pd.to_numeric(f9[c])

    def folio109(r):
        if r.k == 'E' or r.t == '01':
            m = re.match(r'^Pago:\s*\S+\s+(\d+)', r.g or '')
            if m: return int(m.group(1))
        if r.t == 'OD' and int(r.n) > 0: return int(r.n)
        m = re.search(r'(\d{2,})', r.g or '')
        return int(m.group(1)) if m else None
    f9['fo'] = f9.apply(folio109, axis=1) if len(f9) else None
    fmap = {}
    for r in f9[f9.fo.notna() & (f9.c != '00000000')].sort_values('f').itertuples():
        a = fmap.setdefault((r.e, int(r.fo)), {'s': 0, 'nom': None, 'pg': [], 'ces': []})
        a['s'] += r.h - r.d
        nom = (r.nom or '').strip()
        if r.h > 0: a['nom'] = nom or a['nom']; a['ces'].append(f"{r.c} · {_ddmmyyyy(r.f)} · {nom} · {_fmt(r.h)}")
        if r.k == 'E' and r.d > 0: a['pg'].append(f"{r.c} · {_ddmmyyyy(r.f)} · {nom} · {_fmt(r.d)}")

    # --- notas de crédito aplicadas a facturas: calce de Softland si existe; si no, referencia del DTE
    calces = enlaces_calce(calce)
    nc_estado = {(q.emp, q.aux, q.folio): q.est_sf for q in x[x.tipo == '61'].itertuples()}
    nc_monto = {(q.emp, q.aux, q.folio): abs(q.tot) for q in x[x.tipo == '61'].itertuples()}
    nc_a_fac = {}   # (emp, aux, folio_NC) -> [(folio_factura, monto, origen)]
    for kk, lst in calces.items(): nc_a_fac[kk] = [(f, m, 'calce') for f, m in lst]
    for q in x[x.tipo == '61'].itertuples():
        kk = (q.emp, q.aux, q.folio)
        if kk in nc_a_fac or pd.isna(q.tipo_documento_referencia) or pd.isna(q.folio_documento_referencia): continue
        if int(float(q.tipo_documento_referencia)) in (33, 34):
            nc_a_fac[kk] = [(int(float(q.folio_documento_referencia)), abs(q.tot), 'DTE')]
    fac_nc = {}     # (emp, aux, folio_factura) -> [dict(f, m, o, cont)]
    for (e, aux, fn), lst in nc_a_fac.items():
        for f, m, origen in lst:
            cont = nc_estado.get((e, aux, fn), 'Contabilizada') == 'Contabilizada'   # NC solo de calce (fuera del universo): ya contabilizada
            fac_nc.setdefault((e, aux, int(f)), []).append(dict(f=int(fn), m=float(m), o=origen, c=cont))

    base_doc = {(q.emp, q.aux, q.tipo, int(q.folio)): (abs(q.m) if pd.notna(q.m) else abs(q.tot))
                for q in x[(x.est_sf == 'Contabilizada') & x.tipo.isin(['33', '34'])].itertuples()}
    ncap = {}       # (emp, aux, folio_factura) -> monto neteado por NC contabilizadas (tope: el monto de la factura)
    for (e, aux, fo), lst in fac_nc.items():
        b = base_doc.get((e, aux, '33', fo), base_doc.get((e, aux, '34', fo)))
        if b is None: continue
        tot = 0
        for n in lst:
            n['u'] = round(min(n['m'], max(0, b - tot))) if n['c'] else 0
            tot += n['u']
        if tot: ncap[(e, aux, fo)] = tot
    # el saldo FIFO se reparte de nuevo entre los documentos de cada proveedor con NC aplicadas, descontando lo neteado
    afect = {(e, aux) for (e, aux, fo) in ncap}
    if afect and len(opn):
        fe_doc = {(q.emp, q.aux, q.tipo, int(q.folio)): str(q.fe)[:10] for q in x.itertuples()}
        for (e, aux), grp in opn.groupby(['e', 'aux']):
            if (e, aux) not in afect: continue
            S = grp.abierto.sum(); newer = 0
            gg = grp.assign(fe=[fe_doc.get((e, aux, str(t), int(n)), '1900') for t, n in zip(grp.t, grp.n)]).sort_values(['fe', 'n'], ascending=False)
            for i, q in gg.iterrows():
                rem = max(0, q.amt - q.pe - q.pt - q.pf - ncap.get((e, aux, int(q.n)), 0))
                a = max(0, min(rem, S - newer)); opn.at[i, 'abierto'] = round(a); newer += rem
        opmap = {(q.e, q.aux, str(q.t), int(q.n)): q for q in opn.itertuples()}

    # --- armado del registro por documento
    out = []
    for _, r in x.iterrows():
        sg = -1 if r.tipo == '61' else 1
        o = dict(rs=r.rs, t=TIPO[r.tipo], fo=int(r.folio), pv=(_v(r.razon_social_emisor) or '').strip(), rut=r.tax_id_emisor,
                 fe=_d(r.fe), fi=_d(r.fi), fv=_d(r.fv),
                 neto=sg * round((r.monto_neto if pd.notna(r.monto_neto) else 0) + (r.monto_exento if pd.notna(r.monto_exento) else 0)),
                 iva=sg * round(r.monto_iva if pd.notna(r.monto_iva) else 0), tot=sg * round(r.tot),
                 cg=_v(r.cg), cgc=_v(r.cg_cod),
                 eic=(_v(r.est_ic) or '').replace('Factura Documento ', '') or None,
                 asoc=(_v(r.asoc) or '').replace('Factura Documento ', '') or None,
                 pago=(_v(r.pago) or '').replace('Factura_', '').replace('_', ' ') or None,
                 acept=_v(r.acept), ap=_v(r.aprobadores),
                 ces=_v(r.razon_social_cesionario) if str(r.tiene_cesion).lower() in ('true', '1') else None,
                 ref=(f"{TIPO.get(str(int(float(r.tipo_documento_referencia))), r.tipo_documento_referencia)} {r.folio_documento_referencia}"
                      if pd.notna(r.tipo_documento_referencia) and r.tipo in ('56', '61') else None),
                 oc=_v(r.ocs),
                 refp=(f"{REF_SII.get(str(int(r.tipo_documento_referencia)), 'Ref. ' + str(int(r.tipo_documento_referencia)))} {r.folio_documento_referencia}"
                       if pd.notna(r.tipo_documento_referencia) and r.tipo in ('33', '34') and pd.notna(r.folio_documento_referencia) else None),
                 sf=r.est_sf, cpe=(_v(r.s_np) if r.est_sf == 'Provisoria / nula' else None), cpb=_v(r.c), fc=_d(r.f),
                 cta=(f"{r.gc} {str(r.gd).strip()}" if pd.notna(r.gc) else None),
                 cc=(str(r.ccd).strip() if pd.notna(r.ccd) else None), emp=_v(r.emp),
                 msf=(round(abs(r.m)) if pd.notna(r.m) and abs(abs(r.m) - r.tot) > TOL else None))
        ss = bool(r.solosf == True)   # noqa: E712  (NaN en filas de iConstruye)
        es = est_sii.get((r.tax_id_emisor, int(r.tipo), int(r.folio), r.tax_id_receptor))
        if ss:
            o['eic'] = ETQ_SOLO_SF; o['solosf'] = 1
            for kk in ('neto', 'iva', 'fi', 'fv'): o.pop(kk, None)
            if r.tipo in ('33', '34'): o['sii'] = SII_HIST
        elif r.tipo in ('33', '34', '43'):
            o.update(es or dict(sii='Sin consulta'))
            # Ley 19.983: sin reclamo dentro de 8 días corridos desde la recepción, la factura se da por aceptada
            if o['sii'] == 'Sin pronunciamiento':
                rec = _d(r.fi) or _d(r.fe)
                if rec and (date.fromisoformat(corte) - date.fromisoformat(rec)).days > PLAZO_RECLAMO_DIAS:
                    o['sii'] = 'Aceptada (Aut)'
        # rechazada por regla de aceptación en el SII: no es «pendiente de contabilizar» ni entra a la antigüedad
        if o.get('sii') == 'Rechazada' and r.est_sf == 'Pendiente de contabilizar':
            o['sf'] = 'Rechazada en SII'
        if str(r.pg) == 'SI': o['fpi'] = _d(r.fp)
        o['fep'] = _d(r.fep)

        ncs = fac_nc.get((r.emp, r.aux, int(r.folio))) if r.tipo in ('33', '34') else None
        if ncs:
            o['nc'] = [f"{n['f']} · -{_fmt(n['m'])}" + ('' if n['c'] else ' · sin contabilizar') + (' · calce' if n['o'] == 'calce' else ' · ref. DTE') for n in ncs]
            if ncap.get((r.emp, r.aux, int(r.folio))): o['ncm'] = ncap[(r.emp, r.aux, int(r.folio))]
        if r.tipo == '61':
            ps = 'Nota de crédito'
        elif r.est_sf == 'Pendiente de contabilizar' and o.get('sii') == 'Rechazada':
            ps = 'Rechazada en SII'
        elif r.est_sf != 'Contabilizada':
            ps = 'Sin contabilizar'
        else:
            base = abs(r.m) if pd.notna(r.m) else abs(r.tot)
            key = (r.emp, r.aux, r.tipo, int(r.folio))
            nc_ap = ncap.get((r.emp, r.aux, int(r.folio)), 0) if r.tipo in ('33', '34') else 0
            a = agg.get(key)
            if a:
                if a['eg']: o['eg'] = a['eg']
                if a['tr']: o['trs'] = a['tr']
            q = opmap.get(key)
            e = q.pe if q is not None else 0
            tt = q.pt if q is not None else 0
            pf = q.pf if q is not None else 0
            ab = q.abierto if q is not None else base
            if nc_ap:                       # la NC rebaja lo que se debe a esta factura: no se asigna saldo por FIFO
                base_ncx = base; base = max(0, base - nc_ap); ab = min(ab, max(0, base - e - tt)) if base > TOL else 0
            fa = fmap.get((r.emp, int(r.folio))) if pf > 0 else None
            if fa:
                o['fact'] = fa['nom']; o['fces'] = fa['ces']
                if fa['pg']: o['fpag'] = fa['pg']
            if nc_ap and base <= TOL and e + tt <= TOL:
                ps = 'Neteada por NC'
            elif pf > 0 and e + tt < base - TOL:            # cedida a factoring
                fs = max(0, fa['s']) if fa else pf
                if fs <= TOL: ps = 'Pagada a factoring'
                else: ps = 'En Factoring'; o['saldo'] = round(fs)
            elif e + tt >= base - TOL:
                ps = 'Pagada (egreso)' if e > 0 else 'Cancelada por traspaso'
            elif ab <= TOL:
                ps = 'Pagada sin referencia'
            elif ab < base - TOL:
                ps = 'Pago parcial'; o['saldo'] = round(ab)
            else:
                ps = 'En Proveedores'; o['saldo'] = round(ab)
            o['pagado'] = 0 if ps == 'Neteada por NC' else round(base - o.get('saldo', 0))
        o['ps'] = ps
        if r.tipo == '61':     # a qué factura se aplicó esta nota de crédito
            ap = nc_a_fac.get((r.emp, r.aux, int(r.folio)))
            if ap: o['ncap'] = [f"{f} · {_fmt(m)}" + (' · calce' if orig == 'calce' else ' · ref. DTE') for f, m, orig in ap]
        if o.get('ces') and ps not in ('En Factoring', 'Pagada a factoring', 'Nota de crédito'):
            o['cesx'] = 1   # cedida en SII pero no traspasada a Factoring en Softland
        out.append({kk: vv for kk, vv in o.items() if vv is not None})
    return out


# ------------------------------------------------------------------ salidas
def excel(out, ruta):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    cols = [('rs', 'Razón social'), ('t', 'Tipo'), ('fo', 'Folio'), ('pv', 'Proveedor'), ('rut', 'RUT proveedor'),
            ('fe', 'Emisión'), ('fi', 'Recepción IConstruye'), ('fv', 'Vencimiento'), ('neto', 'Neto/exento'), ('iva', 'IVA'),
            ('tot', 'Total'), ('cg', 'Centro de gestión'), ('eic', 'Estado IConstruye'), ('pago', 'Pago IConstruye'),
            ('fpi', 'Fecha pago IC'), ('fep', 'Fecha estimada pago IC'), ('sf', 'Estado Softland'), ('emp', 'Empresa'),
            ('cpb', 'Comprobante'), ('fc', 'Fecha contab.'), ('cta', 'Cuenta'), ('cc', 'Centro de costo'),
            ('ps', 'Pago Softland'), ('pagado', 'Pagado'), ('saldo', 'Saldo'), ('fact', 'Factoring'), ('ces', 'Cesión SII'),
            ('eg', 'Egresos'), ('fpag', 'Pagos a factoring'), ('ap', 'Aprobadores IC'),
            ('oc', 'OC'), ('refp', 'Referencia del proveedor'), ('sii', 'Estado SII'), ('siim', 'Motivo rechazo SII'),
            ('siif', 'Fecha evento SII'), ('siic', 'Última consulta SII'),
            ('nc', 'Notas de crédito'), ('ncm', 'Monto neteado por NC'), ('ncap', 'NC aplicada a factura')]
    wb = Workbook(); ws = wb.active; ws.title = 'Detalle'
    F = Font(name='Calibri', size=10)
    ws.append([c[1].upper() for c in cols])
    for c in ws[1]:
        c.font = Font(name='Calibri', size=10, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='111111')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for o in out:
        row = []
        for kk, _ in cols:
            v = o.get(kk)
            if isinstance(v, list): v = ' | '.join(v)
            if kk == 'nc' and v is None and o['t'] in ('Factura', 'Factura exenta'): v = 'Sin NC'
            if kk in ('fe', 'fi', 'fv', 'fpi', 'fep', 'fc', 'siif') and v: v = date.fromisoformat(v)
            row.append(v)
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = F
            if isinstance(c.value, (int, float)) and c.column_letter != 'C': c.number_format = '#,##0'
            if isinstance(c.value, date): c.number_format = 'DD-MM-YYYY'
    n = ws.max_row
    ws.append(['TOTAL'] + [None] * 9 + [f'=SUM(K2:K{n})'])
    for c in ws[n + 1]:
        c.font = Font(name='Calibri', size=10, bold=True); c.fill = PatternFill('solid', fgColor='F2F2F3')
        c.border = Border(top=Side(style='medium', color='111111')); c.number_format = '#,##0'
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f'A1:{ws.cell(1, len(cols)).column_letter}{n}'
    for i in range(1, len(cols) + 1): ws.column_dimensions[ws.cell(1, i).column_letter].width = 16
    for l, w in (('A', 30), ('D', 32), ('L', 28), ('U', 34), ('AB', 50), ('AD', 50)): ws.column_dimensions[l].width = w
    wb.save(ruta)


VARIANTES = {  # nombre de archivo → (alcance, título, marca, h1, universo)
    'Facturas_Recibidas.html': (None, 'Facturas Recibidas {anio}', 'Trei Inmobiliaria · VCB · IConstruye ↔ Softland',
                                'Facturas recibidas y pago — {anio}', 'para las 10 razones sociales del grupo'),
    'Facturas_Recibidas_VCB.html': ('VCB Constructora SpA', 'Facturas Recibidas {anio} — VCB Constructora SpA',
                                    'VCB Constructora · IConstruye ↔ Softland',
                                    'VCB Constructora SpA — facturas recibidas y pago {anio}', 'para VCB Constructora SpA'),
}


def _ddmm(iso):
    return f"{iso[8:10]}-{iso[5:7]}-{iso[:4]}" if iso else '—'


def armar_meta(out, a, ahora):
    ult = lambda k: max((o[k] for o in out if o.get(k)), default=None)
    sc = [o['siic'] for o in out if o.get('siic')]
    return dict(generado=ahora.strftime('%Y-%m-%dT%H:%M'), desde=a.desde, hasta=a.hasta, corte=a.corte,
                ic_ultimo=ult('fi'), sf_ultimo=ult('fc'), sii_ultimo=max(sc) if sc else None)


def render(plantilla, out, meta, nombre):
    scope, titulo, marca, h1, univ = VARIANTES[nombre]
    anio = meta['desde'][:4]
    filas = [o for o in out if scope is None or o['rs'] == scope]
    js = lambda v: json.dumps(v, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
    html = open(plantilla, encoding='utf-8').read()
    for k, v in (('__TITULO__', titulo.format(anio=anio)), ('__MARCA__', marca), ('__H1__', h1.format(anio=anio)), ('__UNIVERSO__', univ)):
        html = html.replace(k, v)
    return (html.replace('__SCOPE__', js(scope)).replace('__META__', js(meta)).replace('__DATA__', js(filas)))


def main():
    hoy = datetime.now(ZoneInfo('America/Santiago'))
    ap = argparse.ArgumentParser(description='Cruce facturas IConstruye ↔ Softland')
    ap.add_argument('--desde', default=f'{hoy.year}-01-01', help='fecha de emisión desde (incluida); por defecto 1-ene del año en curso')
    ap.add_argument('--hasta', default=str((hoy + timedelta(days=1)).date()), help='fecha de emisión hasta (excluida); por defecto mañana, es decir hasta hoy inclusive')
    ap.add_argument('--corte', default=str(hoy.date()), help='fecha de corte para vencidos; por defecto hoy')
    ap.add_argument('--plantilla', help='HTML del informe con los marcadores __DATA__, __SCOPE__, __META__, __TITULO__, __MARCA__, __H1__, __UNIVERSO__')
    ap.add_argument('--salida', default='.', help='carpeta de salida')
    ap.add_argument('--silencioso', action='store_true', help='no imprime cifras del informe (para ejecuciones con registro público, como GitHub Actions)')
    a = ap.parse_args()
    os.makedirs(a.salida, exist_ok=True)

    ic, sf, pay, opn, f9, sii, calce = extraer(a)
    out = cruzar(ic, sf, pay, opn, f9, a.corte, sii, a.desde[:4], calce)
    meta = armar_meta(out, a, hoy)

    json.dump(dict(meta=meta, datos=out), open(os.path.join(a.salida, 'facturas_cruce.json'), 'w', encoding='utf-8'),
              ensure_ascii=False, separators=(',', ':'))
    excel(out, os.path.join(a.salida, 'Facturas_IConstruye_Softland.xlsx'))
    if a.plantilla:
        for nombre in VARIANTES:
            open(os.path.join(a.salida, nombre), 'w', encoding='utf-8', newline='\n').write(render(a.plantilla, out, meta, nombre))
            print('Escrito', nombre)

    if a.silencioso:
        print('Informe generado'); return
    from collections import Counter
    print(f"{len(out)} documentos · {a.desde} → {a.corte}")
    print('Contabilización:', dict(Counter(o['sf'] for o in out)))
    print('Pago Softland:  ', dict(Counter(o['ps'] for o in out)))
    print('Estado SII:     ', dict(Counter(o.get('sii', '—') for o in out)))


if __name__ == '__main__':
    main()
